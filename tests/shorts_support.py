"""Synthetic fixtures for the net short position connectors. Holders and issuers are invented names."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from datetime import UTC, date, datetime
from typing import Any
from xml.sax.saxutils import escape

import httpx
from conftest import make_isin, make_lei

from disclosure_atlas.connectors.base import RawDocument
from disclosure_atlas.http import PoliteClient

FETCHED_AT = datetime(2026, 10, 20, 6, 0, tzinfo=UTC)
HOLDER_LEI = make_lei("549300EXAMPLEHOLD1")
ISSUER_LEI = make_lei("969500EXAMPLEISSU1")


def isin(country: str) -> str:
    return make_isin(f"{country}000EXAMPL")


def doc(source: str, content: bytes, suffix: str, window: tuple[date, date] | None = None, **meta: Any) -> RawDocument:
    if window:
        meta["window"] = [window[0].isoformat(), window[1].isoformat()]
    return RawDocument(source=source, url="https://example.invalid/", params={}, content=content, content_type=None,
                       fetched_at=FETCHED_AT, suffix=suffix, meta=meta)


def ods_bytes(rows: list[list[Any]], name: str = "Blad1") -> bytes:
    """A minimal OpenDocument spreadsheet, with the repeated empty trailing cells real exports carry."""
    def cell(value: Any) -> str:
        if value is None:
            return "<table:table-cell/>"
        if isinstance(value, float):
            return (f'<table:table-cell office:value-type="float" office:value="{value}">'
                    f"<text:p>{str(value).replace('.', ',')}</text:p></table:table-cell>")
        if isinstance(value, date):
            return (f'<table:table-cell office:value-type="date" office:date-value="{value.isoformat()}T00:00:00">'
                    f"<text:p>{value.isoformat()}</text:p></table:table-cell>")
        lines = "".join(f"<text:p>{escape(line)}</text:p>" for line in str(value).split("\n"))
        return f'<table:table-cell office:value-type="string">{lines}</table:table-cell>'

    body = "".join(
        "<table:table-row>" + "".join(cell(v) for v in row)
        + '<table:table-cell table:number-columns-repeated="16378"/></table:table-row>'
        for row in rows
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        f'<office:body><office:spreadsheet><table:table table:name="{name}">{body}</table:table>'
        "</office:spreadsheet></office:body></office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", content)
    return buffer.getvalue()


def xlsx_bytes(sheets: dict[str, list[list[Any]]]) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def mock_client(handler: Callable[[httpx.Request], httpx.Response]) -> PoliteClient:
    client = PoliteClient(min_interval=0, retries=0)
    client._client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=True)
    return client
